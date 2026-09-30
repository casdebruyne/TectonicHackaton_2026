import base64
import io
import json
import sqlite3
from datetime import datetime

import docx
import openpyxl
import pandas as pd
import pypdf
import streamlit as st
from PIL import ExifTags, Image

st.set_page_config(page_title="Knowledge Reliability", page_icon="🛡️", layout="wide")

DB = "kennis.db"
DREMPEL = 75  # from this score onwards a document counts as "reliable"

# Resource limits to prevent unbounded storage and cross-user resource exhaustion
MAX_FILES_PER_UPLOAD = 20  # Maximum files per upload batch
MAX_FILE_SIZE_MB = 10  # Maximum size per file in MB
MAX_BATCH_SIZE_MB = 50  # Maximum total size per upload batch in MB
MAX_TOTAL_DOCUMENTS = 1000  # Maximum total documents in database
MAX_DATABASE_SIZE_MB = 500  # Maximum total database size in MB

# Maximum points per factor (total 100)
GEWICHTEN = {
    "Recency": 10,
    "Author": 30,
    "AI percentage": 30,
    "Metadata present": 15,
    "File type": 15,
}
TYPE_SCORE = {"pdf": 1.0, "docx": 0.9, "xlsx": 0.8, "jpg": 0.4, "jpeg": 0.4, "png": 0.4}
MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
}
TALEN = ["Dutch", "French", "English", "German", "Other"]
LANDEN = ["Belgium", "Netherlands", "France", "Germany", "United Kingdom", "Other"]

# AI content levels (unknown = "I don't know" checkbox, not counted in the score)
AI_NIVEAUS = {
    "0% · No AI": 0,
    "25% · Little AI": 25,
    "50% · Half": 50,
    "75% · Mostly AI": 75,
    "100% · Fully AI": 100,
}


def ai_tekst(ai):
    return "Unknown" if ai is None or pd.isna(ai) else f"{int(ai)}%"


def sanitize_csv_formula(value):
    """
    Neutralizes potential spreadsheet formula injection by prefixing
    formula-leading characters with a single quote. This prevents
    interpretation as a formula when the CSV is opened in spreadsheet software.
    
    Formula markers: = + - @ (and tab/carriage return which are less common)
    """
    if not value or not isinstance(value, str):
        return value
    
    # Strip whitespace first
    value = value.strip()
    
    # Check if the value starts with a formula marker
    if value and value[0] in ('=', '+', '-', '@', '\t', '\r'):
        # Prefix with single quote to neutralize the formula
        return "'" + value
    
    return value


# ---------------------------------------------------------------
# Style
# ---------------------------------------------------------------
st.markdown(
    """
<style>
section[data-testid="stSidebar"] { border-right: 1px solid rgba(128,128,128,.2); }
section[data-testid="stSidebar"] .stButton > button {
    justify-content: flex-start; text-align: left; border-radius: 12px;
    padding: .65rem 1rem; font-weight: 600; margin-bottom: .15rem;
}
.brand { font-size: 1.5rem; font-weight: 800; line-height: 1.1;
    background: linear-gradient(90deg,#FF4B4B,#FF8C42); -webkit-background-clip: text;
    -webkit-text-fill-color: transparent; }
.brand-sub { opacity: .6; font-size: .85rem; margin-bottom: .8rem; }
</style>
""",
    unsafe_allow_html=True,
)


# ---------------------------------------------------------------
# Database
# ---------------------------------------------------------------
def get_con():
    con = sqlite3.connect(DB)
    con.execute(
        """CREATE TABLE IF NOT EXISTS documenten (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            bestandsnaam TEXT, onderwerp TEXT, taal TEXT, auteur TEXT, land TEXT,
            ai_percentage INTEGER, bestandsdatum TEXT, toegevoegd_op TEXT,
            score INTEGER, uitleg TEXT, inhoud BLOB)"""
    )
    kolommen = [r[1] for r in con.execute("PRAGMA table_info(documenten)")]
    if "inhoud" not in kolommen:  # update old database
        con.execute("ALTER TABLE documenten ADD COLUMN inhoud BLOB")
    return con


def lees_database():
    con = get_con()
    df = pd.read_sql(
        """SELECT id, bestandsnaam, onderwerp, taal, auteur, land, ai_percentage,
                  bestandsdatum, toegevoegd_op, score, uitleg
           FROM documenten ORDER BY score DESC""",
        con,
    )
    con.close()
    return df


def get_database_stats():
    """Returns current database statistics for quota enforcement."""
    con = get_con()
    # Get document count
    doc_count = con.execute("SELECT COUNT(*) FROM documenten").fetchone()[0]
    # Get total size of stored BLOBs
    total_size = con.execute("SELECT SUM(LENGTH(inhoud)) FROM documenten WHERE inhoud IS NOT NULL").fetchone()[0] or 0
    con.close()
    return doc_count, total_size


def check_upload_quota(new_files_data):
    """
    Validates that adding new files won't exceed resource quotas.
    Returns (is_valid, error_message).
    """
    # Check number of files in this batch
    if len(new_files_data) > MAX_FILES_PER_UPLOAD:
        return False, f"Too many files at once. Maximum is {MAX_FILES_PER_UPLOAD} files per upload."
    
    # Check individual file sizes and batch total
    batch_size = 0
    for data in new_files_data:
        file_size_mb = len(data) / (1024 * 1024)
        if file_size_mb > MAX_FILE_SIZE_MB:
            return False, f"A file is too large ({file_size_mb:.1f} MB). Maximum is {MAX_FILE_SIZE_MB} MB per file."
        batch_size += len(data)
    
    batch_size_mb = batch_size / (1024 * 1024)
    if batch_size_mb > MAX_BATCH_SIZE_MB:
        return False, f"Total upload size ({batch_size_mb:.1f} MB) exceeds the maximum of {MAX_BATCH_SIZE_MB} MB."
    
    # Check database limits
    doc_count, current_db_size = get_database_stats()
    
    if doc_count + len(new_files_data) > MAX_TOTAL_DOCUMENTS:
        return False, f"Database limit reached. Maximum number of documents is {MAX_TOTAL_DOCUMENTS}. Current: {doc_count}."
    
    projected_size_mb = (current_db_size + batch_size) / (1024 * 1024)
    if projected_size_mb > MAX_DATABASE_SIZE_MB:
        return False, f"Database storage limit reached. Maximum is {MAX_DATABASE_SIZE_MB} MB. Current: {current_db_size / (1024 * 1024):.1f} MB."
    
    return True, None


def lees_inhoud(doc_id):
    con = get_con()
    rij = con.execute("SELECT inhoud FROM documenten WHERE id=?", (int(doc_id),)).fetchone()
    con.close()
    return rij[0] if rij else None


def aantal_betrouwbaar(auteur, uitsluiten_id=None):
    """Number of documents by this author with score >= DREMPEL (own document excluded)."""
    if not auteur or not auteur.strip():
        return 0
    con = get_con()
    n = con.execute(
        "SELECT COUNT(*) FROM documenten WHERE lower(trim(auteur)) = ? AND score >= ? AND id != ?",
        (auteur.strip().lower(), DREMPEL, uitsluiten_id if uitsluiten_id is not None else -1),
    ).fetchone()[0]
    con.close()
    return n


# ---------------------------------------------------------------
# Metadata from the file itself
# ---------------------------------------------------------------
def lees_metadata(naam, data):
    naam = naam.lower()
    try:
        if naam.endswith(".pdf"):
            m = pypdf.PdfReader(io.BytesIO(data)).metadata
            if m:
                d = m.modification_date
                return (d.replace(tzinfo=None) if d else None), m.author
        elif naam.endswith(".docx"):
            p = docx.Document(io.BytesIO(data)).core_properties
            return p.modified, (p.last_modified_by or p.author)
        elif naam.endswith(".xlsx"):
            p = openpyxl.load_workbook(io.BytesIO(data), read_only=True).properties
            return p.modified, (p.lastModifiedBy or p.creator)
        elif naam.endswith((".jpg", ".jpeg", ".png")):
            exif = Image.open(io.BytesIO(data))._getexif()
            if exif:
                for tag_id, waarde in exif.items():
                    if ExifTags.TAGS.get(tag_id) in ("DateTimeOriginal", "DateTime"):
                        return datetime.strptime(waarde, "%Y:%m:%d %H:%M:%S"), None
    except Exception:
        pass
    return None, None


# ---------------------------------------------------------------
# Score (1-100)
# ---------------------------------------------------------------
def parse_upload(tekst):
    try:
        return datetime.strptime(tekst, "%Y-%m-%d %H:%M") if tekst else None
    except ValueError:
        return None


def bereken_score(d, uitsluiten_id=None):
    f = {}

    # Recency: declines slowly (halves every 5 years) and never drops below 40%.
    # Without a file date we use the upload date, but that says less (max. 50%).
    if d["datum"]:
        jaren = (datetime.now() - d["datum"]).days / 365
        f["Recency"] = max(0.4, 0.5 ** (jaren / 5))
    elif d.get("upload"):
        jaren = (datetime.now() - d["upload"]).days / 365
        f["Recency"] = min(0.5, max(0.4, 0.5 ** (jaren / 5)))
    else:
        f["Recency"] = 0.3

    # Author: based on the number of reliable documents by this author in the database
    # New authors without a track record receive 0 points to prevent score manipulation
    if not d["auteur"].strip():
        f["Author"] = 0.0
    else:
        n = aantal_betrouwbaar(d["auteur"], uitsluiten_id)
        if n == 0:
            f["Author"] = 0.0  # new author without track record: 0%
        else:
            f["Author"] = min(1.0, n / 5)  # from 1 reliable doc: 20%, from 5: 100%

    # AI percentage: unknown (None) = factor is not counted
    if d["ai"] is not None:
        f["AI percentage"] = 1 - d["ai"] / 100

    f["Metadata present"] = 1.0 if d["datum"] else 0.0
    f["File type"] = TYPE_SCORE.get(d["bestandsnaam"].rsplit(".", 1)[-1].lower(), 0.3)

    # Only the factors that count; the score is then scaled to 100
    max_totaal = sum(GEWICHTEN[k] for k in f)
    punten = {k: round(f[k] * GEWICHTEN[k] * 100 / max_totaal, 1) for k in f}
    score = max(1, min(100, round(sum(punten.values()))))
    return score, punten


def rating(score):
    if score >= DREMPEL:
        return "🟢 Reliable"
    if score >= 50:
        return "🟠 Caution"
    return "🔴 Do not trust"


def herbereken_alles(max_rondes=5):
    """Recalculates all stored scores, until nothing changes anymore."""
    con = get_con()
    rijen = con.execute(
        "SELECT id, bestandsnaam, auteur, ai_percentage, bestandsdatum, toegevoegd_op, score FROM documenten"
    ).fetchall()
    con.close()

    huidige = {r[0]: r[6] for r in rijen}

    for _ in range(max_rondes):
        veranderd = False
        for doc_id, naam, auteur, ai, datum, upload, _oud in rijen:
            d = {
                "bestandsnaam": naam,
                "auteur": auteur or "",
                "ai": ai,  # None stays None (unknown)
                "datum": datetime.strptime(datum, "%Y-%m-%d") if datum else None,
                "upload": parse_upload(upload),
            }
            score, punten = bereken_score(d, doc_id)
            con = get_con()
            con.execute(
                "UPDATE documenten SET score=?, uitleg=? WHERE id=?",
                (score, json.dumps(punten), doc_id),
            )
            con.commit()
            con.close()
            if score != huidige[doc_id]:
                huidige[doc_id] = score
                veranderd = True
        if not veranderd:
            break


# ---------------------------------------------------------------
# Open document (preview + download)
# ---------------------------------------------------------------
def toon_document(naam, data):
    ext = naam.rsplit(".", 1)[-1].lower()
    st.download_button(
        f"⬇️ Open / download {naam}", data, file_name=naam,
        mime=MIME.get(ext, "application/octet-stream"), type="primary",
    )
    if ext == "pdf":
        b64 = base64.b64encode(data).decode()
        st.markdown(
            f'<iframe src="data:application/pdf;base64,{b64}" width="100%" height="700"></iframe>',
            unsafe_allow_html=True,
        )
    elif ext in ("jpg", "jpeg", "png"):
        st.image(data, use_container_width=True)
    elif ext == "docx":
        tekst = "\n\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs if p.text.strip())
        st.text_area("Content", tekst or "(no text found)", height=400)
    elif ext == "xlsx":
        for blad, tabel in pd.read_excel(io.BytesIO(data), sheet_name=None).items():
            st.caption(f"Sheet: {blad}")
            st.dataframe(tabel, use_container_width=True)


# ---------------------------------------------------------------
# Sidebar / navigation
# ---------------------------------------------------------------
st.session_state.setdefault("pagina", "upload")
st.session_state.setdefault("wachtrij", [])
st.session_state.setdefault("up_nr", 0)

df_all = lees_database()
n_wacht = len(st.session_state.wachtrij)

with st.sidebar:
    st.markdown('<div class="brand">🛡️ Knowledge<br>Reliability</div>', unsafe_allow_html=True)
    st.markdown('<div class="brand-sub">Find it. Understand it. Trust it.</div>', unsafe_allow_html=True)

    menu = [
        ("upload", "📤", "1 · Upload"),
        ("beoordeel", "🧮", "2 · Assess" + (f"  ({n_wacht})" if n_wacht else "")),
        ("zoek", "🔎", "3 · Search"),
    ]
    for sleutel, icoon, label in menu:
        actief = st.session_state.pagina == sleutel
        if st.button(f"{icoon}  {label}", key=f"nav_{sleutel}", use_container_width=True,
                     type="primary" if actief else "secondary"):
            st.session_state.pagina = sleutel
            st.rerun()

    st.divider()
    s1, s2 = st.columns(2)
    s1.metric("In database", len(df_all))
    s2.metric("Avg. score", f"{df_all['score'].mean():.0f}" if len(df_all) else "–")
    st.caption(f"⏳ {n_wacht} document(s) in queue")
    st.caption(f"🟢 Reliable = score ≥ {DREMPEL}")

pagina = st.session_state.pagina


def ga_naar(p):
    st.session_state.pagina = p


# ===============================================================
# PAGE 1: Upload
# ===============================================================
if pagina == "upload":
    st.title("📤 Upload documents")
    st.write("Upload documents and fill in the details for each document.")

    bestanden = st.file_uploader(
        "Choose or drop your files here (PDF, Word, Excel, images)",
        accept_multiple_files=True,
        key=f"up_{st.session_state.up_nr}",
    )

    invoer = []
    for i, f in enumerate(bestanden or []):
        with st.expander(f"📄 {f.name}", expanded=True):
            c1, c2 = st.columns(2)
            onderwerp = c1.text_input("Subject", key=f"onderwerp_{i}_{f.name}")
            # Extract metadata early to show detected author (read-only)
            data = f.getvalue()
            datum, meta_auteur = lees_metadata(f.name, data)
            if meta_auteur and meta_auteur.strip():
                c2.text_input("Author (from file metadata)", value=meta_auteur, disabled=True, key=f"auteur_{i}_{f.name}")
            else:
                c2.caption("⚠️ No author found in file metadata")
            c3, c4 = st.columns(2)
            taal = c3.selectbox("Language", TALEN, key=f"taal_{i}_{f.name}")
            land = c4.selectbox("Country", LANDEN, key=f"land_{i}_{f.name}")
            # AI percentage is always unknown - cannot be user-specified to prevent score manipulation
            ai = None
            st.caption("ℹ️ AI content is not included in the score (automatic detection not yet available).")
            invoer.append((f, onderwerp, meta_auteur, taal, land, ai, datum, data))

    if invoer:
        if st.button("➡️ Send to assessment", type="primary"):
            ontbreekt = [f.name for f, o, *_ in invoer if not o.strip()]
            if ontbreekt:
                st.error("Fill in the subject for: " + ", ".join(ontbreekt))
            else:
                # Validate upload quota before adding to queue
                new_files_data = [data for f, *_, data in invoer]
                is_valid, error_msg = check_upload_quota(new_files_data)
                
                if not is_valid:
                    st.error(f"❌ Upload rejected: {error_msg}")
                else:
                    for f, onderwerp, auteur, taal, land, ai, datum, data in invoer:
                        # Auteur and AI are now derived from metadata/system, not user input
                        # This prevents score manipulation via forged author names or AI percentages
                        # Sanitize user-controlled fields to prevent CSV formula injection
                        st.session_state.wachtrij.append({
                            "bestandsnaam": sanitize_csv_formula(f.name), 
                            "onderwerp": sanitize_csv_formula(onderwerp.strip()), 
                            "auteur": sanitize_csv_formula(auteur or ""),
                            "taal": taal, "land": land, "ai": ai, "datum": datum,
                            "meta_auteur": auteur, "inhoud": data,
                            "upload": datetime.now(),
                        })
                    st.session_state.up_nr += 1  # clear the uploader
                    st.session_state.pagina = "beoordeel"
                    st.rerun()
    else:
        st.info("Upload at least one file to get started.")


# ===============================================================
# PAGE 2: Assess & save
# ===============================================================
elif pagina == "beoordeel":
    st.title("🧮 Assess & save")
    wachtrij = st.session_state.wachtrij

    if not wachtrij:
        st.info("No documents are ready. Upload documents on page 1 first.")
        st.button("📤 Go to upload", on_click=ga_naar, args=("upload",))
    else:
        resultaten = [(d, *bereken_score(d)) for d in wachtrij]

        overzicht = pd.DataFrame([
            {
                "Document": d["bestandsnaam"],
                "Subject": d["onderwerp"],
                "Author": d["auteur"] or "⚠️ unknown",
                "Author's reliable docs": aantal_betrouwbaar(d["auteur"]),
                "File date": d["datum"].strftime("%Y-%m-%d") if d["datum"] else "Unknown",
                "Uploaded on": d["upload"].strftime("%Y-%m-%d %H:%M"),
                "AI %": ai_tekst(d["ai"]),
                "Score": score,
                "Rating": rating(score),
            }
            for d, score, _ in resultaten
        ])
        st.dataframe(
            overzicht, hide_index=True, use_container_width=True,
            column_config={"Score": st.column_config.ProgressColumn("Score", min_value=0, max_value=100, format="%d")},
        )

        st.subheader("🔍 Why this score?")
        keuze = st.selectbox("Choose a document", [d["bestandsnaam"] for d, _, _ in resultaten])
        d, score, punten = next(r for r in resultaten if r[0]["bestandsnaam"] == keuze)
        c1, c2 = st.columns([1, 2])
        c1.metric("Reliability", f"{score} / 100")
        c1.write(rating(score))
        if d["ai"] is None:
            c1.caption("ℹ️ AI content unknown: not counted, so the other factors weigh more heavily.")
        c2.bar_chart(pd.Series(punten, name="Points"))

        b1, b2 = st.columns(2)
        if b1.button("💾 Save to database", type="primary", use_container_width=True):
            # Re-validate quota at save time to prevent session manipulation or race conditions
            files_data = [d["inhoud"] for d in wachtrij]
            is_valid, error_msg = check_upload_quota(files_data)
            
            if not is_valid:
                st.error(f"❌ Save rejected: {error_msg}")
                st.warning("The database limits have been reached. Delete old documents or contact the administrator.")
            else:
                con = get_con()
                for d, score, punten in resultaten:
                    con.execute(
                        """INSERT INTO documenten
                        (bestandsnaam, onderwerp, taal, auteur, land, ai_percentage,
                         bestandsdatum, toegevoegd_op, score, uitleg, inhoud)
                        VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                        (
                            d["bestandsnaam"], d["onderwerp"], d["taal"], d["auteur"], d["land"], d["ai"],
                            d["datum"].strftime("%Y-%m-%d") if d["datum"] else None,
                            d["upload"].strftime("%Y-%m-%d %H:%M"),
                            score, json.dumps(punten), d["inhoud"],
                        ),
                    )
                con.commit()
                con.close()
                herbereken_alles()  # older documents by the same author move up/down too
                st.session_state.wachtrij = []
                st.session_state.pagina = "zoek"
                st.balloons()
                st.rerun()
        if b2.button("🗑️ Clear queue", use_container_width=True):
            st.session_state.wachtrij = []
            st.rerun()


# ===============================================================
# PAGE 3: Search + open document
# ===============================================================
else:
    st.title("🔎 Search the knowledge base")
    df = df_all

    if df.empty:
        st.info("The database is still empty. Add documents first via pages 1 and 2.")
    else:
        zoek = st.text_input("Search by subject, file name or author")
        c1, c2, c3 = st.columns(3)
        talen = c1.multiselect("Language", sorted(df["taal"].unique()))
        landen = c2.multiselect("Country", sorted(df["land"].unique()))
        min_score = c3.slider("Minimum score", 1, 100, 1)

        res = df.copy()
        if zoek:
            z = zoek.lower()
            res = res[
                res["onderwerp"].str.lower().str.contains(z, na=False)
                | res["bestandsnaam"].str.lower().str.contains(z, na=False)
                | res["auteur"].str.lower().str.contains(z, na=False)
            ]
        if talen:
            res = res[res["taal"].isin(talen)]
        if landen:
            res = res[res["land"].isin(landen)]
        res = res[res["score"] >= min_score].sort_values("score", ascending=False).reset_index(drop=True)
        res["Rating"] = res["score"].apply(rating)
        res["AI %"] = res["ai_percentage"].apply(ai_tekst)

        st.write(f"**{len(res)}** result(s)  ·  👆 click a row to open the document")
        event = st.dataframe(
            res[["bestandsnaam", "onderwerp", "auteur", "taal", "land", "AI %", "toegevoegd_op", "score", "Rating"]],
            hide_index=True, use_container_width=True,
            on_select="rerun", selection_mode="single-row", key="zoektabel",
            column_config={
                "bestandsnaam": "File name",
                "onderwerp": "Subject",
                "auteur": "Author",
                "taal": "Language",
                "land": "Country",
                "score": st.column_config.ProgressColumn("Score", min_value=0, max_value=100, format="%d"),
                "toegevoegd_op": "Uploaded on",
            },
        )

        geselecteerd = event.selection.rows
        if geselecteerd:
            r = res.iloc[geselecteerd[0]]
            st.divider()
            st.subheader(f"📄 {r['bestandsnaam']}")
            k1, k2 = st.columns([1, 2])
            k1.metric("Reliability", f"{r['score']} / 100")
            k1.write(r["Rating"])
            k1.write(f"**Subject:** {r['onderwerp']}")
            k1.write(f"**Author:** {r['auteur'] or '⚠️ unknown'}")
            k1.write(f"**AI content:** {r['AI %']}")
            k1.write(f"**File date:** {r['bestandsdatum'] or 'unknown'}")
            k1.write(f"**Uploaded on:** {r['toegevoegd_op']}")
            k2.caption("How the score is built up")
            k2.bar_chart(pd.Series(json.loads(r["uitleg"]), name="Points"))

            inhoud = lees_inhoud(r["id"])
            if inhoud:
                toon_document(r["bestandsnaam"], inhoud)
            else:
                st.warning("This document was saved before the update and has no file in the database.")
        else:
            st.info("Select a row to see the score explanation and open the document.")

        b1, b2 = st.columns(2)
        if b1.button("🔄 Recalculate all scores", use_container_width=True):
            herbereken_alles()
            st.rerun()
        
        # Sanitize CSV export to prevent formula injection
        # Apply sanitization to user-controlled text fields that could contain formulas
        df_export = df.copy()
        for col in ['bestandsnaam', 'onderwerp', 'auteur']:
            if col in df_export.columns:
                df_export[col] = df_export[col].apply(sanitize_csv_formula)
        
        b2.download_button(
            "⬇️ Download database as CSV",
            df_export.to_csv(index=False).encode("utf-8"),
            file_name="kennisbank.csv", mime="text/csv",
            use_container_width=True,
        )
